# -*- coding: utf-8 -*-
import openpyxl

out_path = "D:/27968/我的文档/Users/27968/Documents/青羊58施工计划项目/上海远香湖_物业备品配件数据说明.xlsx"
src_path = "D:/项目综合资料/天府办公文件/上海远香湖结算单/最总结算表格/02 结算前扣款事项汇总.xlsx"

wb = openpyxl.load_workbook(out_path)
wb_src = openpyxl.load_workbook(src_path, data_only=True)

ws = wb["配品数据总说明"]

# 灯具数据
light_data = {}
for r in range(7, 36):
    n = wb_src["灯具结算"].cell(r, 2).value
    c30 = wb_src["灯具结算"].cell(r, 30).value
    c31 = wb_src["灯具结算"].cell(r, 31).value
    p = wb_src["灯具结算"].cell(r, 6).value
    if n and isinstance(c30, (int,float)):
        light_data[n] = {"ht": c30, "hs": c31 or c30, "p": p}

# 开关面板数据
sw_data = {}
for r in range(6, 39):
    n = wb_src["开关插座面板结算"].cell(r, 2).value
    c29 = wb_src["开关插座面板结算"].cell(r, 29).value
    c30 = wb_src["开关插座面板结算"].cell(r, 30).value
    p = wb_src["开关插座面板结算"].cell(r, 6).value
    if n and isinstance(c29, (int,float)):
        sw_data[n] = {"ht": c29, "hs": c30 or c29, "p": p}

# HDPE管材数据
hdpe_data = {}
for r in range(5, 51):
    n = wb_src["HDPE管材结算"].cell(r, 3).value
    c15 = wb_src["HDPE管材结算"].cell(r, 15).value
    c20 = wb_src["HDPE管材结算"].cell(r, 20).value
    p = wb_src["HDPE管材结算"].cell(r, 9).value
    sp = str(wb_src["HDPE管材结算"].cell(r, 4).value or "")
    if n and isinstance(c20, (int,float)):
        ht = c15 if isinstance(c15, (int,float)) else 0
        hdpe_data[n] = {"ht": ht, "hs": c20, "p": p, "spec": sp}

# 洁具科勒数据
kohler_data = {}
for r in range(5, 10):
    n = wb_src["洁具结算（科勒）"].cell(r, 2).value
    c13 = wb_src["洁具结算（科勒）"].cell(r, 13).value
    p = wb_src["洁具结算（科勒）"].cell(r, 6).value
    if n and isinstance(c13, (int,float)):
        kohler_data[n] = {"ht": c13, "hs": c13, "p": p}

# 户内地漏(HDPE)数据
dl_data = {}
for r in range(5, 9):
    n = wb_src["户内地漏（HDPE）结算"].cell(r, 3).value
    c15 = wb_src["户内地漏（HDPE）结算"].cell(r, 15).value
    c20 = wb_src["户内地漏（HDPE）结算"].cell(r, 20).value
    p = wb_src["户内地漏（HDPE）结算"].cell(r, 9).value
    if n and isinstance(c20, (int,float)):
        ht = c15 if isinstance(c15, (int,float)) else 0
        dl_data[n] = {"ht": ht, "hs": c20, "p": p}

# HDPE电焊管箍数据
gg_data = {}
for r in range(5, 8):
    n = wb_src["HDPE电焊管箍结算"].cell(r, 3).value
    c15 = wb_src["HDPE电焊管箍结算"].cell(r, 15).value
    c20 = wb_src["HDPE电焊管箍结算"].cell(r, 20).value
    p = wb_src["HDPE电焊管箍结算"].cell(r, 9).value
    if n and isinstance(c20, (int,float)):
        ht = c15 if isinstance(c15, (int,float)) else 0
        gg_data[n] = {"ht": ht, "hs": c20, "p": p}

# 浴霸及凉霸
hb_data = {}
for r in range(5, 8):
    n = wb_src["浴霸及凉霸结算"].cell(r, 2).value
    c13 = wb_src["浴霸及凉霸结算"].cell(r, 13).value
    p = wb_src["浴霸及凉霸结算"].cell(r, 6).value
    if n and isinstance(c13, (int,float)):
        hb_data[n] = {"ht": c13, "hs": c13, "p": p}

# 户内配电箱
pd_data = {}
for r in range(5, 8):
    n = wb_src["户内配电箱结算"].cell(r, 2).value
    c13 = wb_src["户内配电箱结算"].cell(r, 13).value
    c14 = wb_src["户内配电箱结算"].cell(r, 14).value
    p = wb_src["户内配电箱结算"].cell(r, 6).value
    if n and isinstance(c14, (int,float)):
        ht = c13 if isinstance(c13, (int,float)) else 0
        pd_data[n] = {"ht": ht, "hs": c14, "p": p}

def find_data(cat, name):
    if cat == "灯具":
        return light_data.get(name)
    elif cat == "开关插座面板":
        return sw_data.get(name)
    elif "HDPE" in str(cat) and "管材" in str(cat):
        return hdpe_data.get(name)
    elif "HDPE" in str(cat) and "地漏" in str(cat):
        return dl_data.get(name)
    elif "HDPE" in str(cat) and "管箍" in str(cat):
        return gg_data.get(name)
    elif "洁具" in str(cat):
        return kohler_data.get(name)
    elif "浴霸" in str(cat):
        return hb_data.get(name)
    elif "配电箱" in str(cat):
        return pd_data.get(name)
    return None

filled = 0
for r in range(6, ws.max_row + 1):
    cat = str(ws.cell(r, 1).value or "")
    name = str(ws.cell(r, 2).value or "")
    
    d = find_data(cat, name)
    if not d:
        # 尝试模糊匹配
        lookup = light_data if cat == "灯具" else sw_data if cat == "开关插座面板" else hdpe_data if "管材" in cat else {}
        for k, v in lookup.items():
            if k and name and (k in name or name in k):
                d = v
                break
    
    if d:
        ht = d["ht"]
        hs = d["hs"]
        pj = round(hs - ht, 2) if isinstance(hs, (int,float)) and isinstance(ht, (int,float)) else 0
        ws.cell(r, 5, ht)
        ws.cell(r, 6, pj)
        filled += 1

# 也对HDPE行按规格型号匹配
for r in range(6, ws.max_row + 1):
    cat = str(ws.cell(r, 1).value or "")
    if "HDPE" in cat and "管材" in cat:
        spec = str(ws.cell(r, 3).value or "")
        for n, d in hdpe_data.items():
            sp = d.get("spec", "")
            if spec and sp and (spec in sp or sp in spec):
                if not ws.cell(r, 5).value:
                    ws.cell(r, 5, d["ht"])
                    ws.cell(r, 6, round(d["hs"] - d["ht"], 2))
                    filled += 1
                break

wb.save(out_path)
print(f"OK: {filled}")
