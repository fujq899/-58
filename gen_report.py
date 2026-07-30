# -*- coding: utf-8 -*-
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

src = r"D:\项目综合资料\天府办公文件\上海远香湖结算单\最总结算表格\02 结算前扣款事项汇总.xlsx"
wb_src = openpyxl.load_workbook(src, data_only=True)

wb = openpyxl.Workbook()

hdr_font = Font(bold=True, size=11, color="FFFFFF")
hdr_fill = PatternFill("solid", fgColor="4472C4")
red_fill = PatternFill("solid", fgColor="FFC7CE")
green_fill = PatternFill("solid", fgColor="C6EFCE")
title_font = Font(bold=True, size=14)
thin_border = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"), bottom=Side(style="thin")
)

def style_header(ws, row, cols):
    for c in range(1, cols+1):
        cell = ws.cell(row=row, column=c)
        cell.font = hdr_font
        cell.fill = hdr_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border

def style_cell(ws, row, cols):
    for c in range(1, cols+1):
        cell = ws.cell(row=row, column=c)
        cell.border = thin_border
        cell.alignment = Alignment(horizontal="center", vertical="center")

# Sheet 1
ws1 = wb.active
ws1.title = "物业配品配件对比总表"
ws1.cell(1,1,"上海嘉定远香湖项目 - 甲供材物业配品配件移交对比表").font = Font(bold=True, size=14, color="4472C4")
ws1.merge_cells("A1:H1")
ws1.cell(2,1,"核销公式：含损耗数量 = 物业配品配件(移交) + 四方单统计(施工领用)").font = Font(size=10, italic=True, color="666666")
ws1.merge_cells("A2:H2")

headers = ["材料类别", "材料名称", "规格型号", "单位", "含损耗数量(合同量)", "四方单统计(施工领用)", "物业配品配件(应交物业)", "配比对账结果"]
for c, h in enumerate(headers, 1):
    ws1.cell(3, c, h)
style_header(ws1, 3, 8)

row = 4

# 灯具结算
ws_light = wb_src["灯具结算"]
for r in range(7, 36):
    name = ws_light.cell(row=r, column=2).value
    unit = ws_light.cell(row=r, column=4).value
    shiji = ws_light.cell(row=r, column=5).value
    hansun = ws_light.cell(row=r, column=31).value
    spec = ws_light.cell(row=r, column=3).value
    if name and isinstance(shiji, (int,float)):
        peijian = (hansun or 0) - shiji if isinstance(hansun, (int,float)) else None
        if peijian is not None:
            ws1.cell(row, 1, "灯具")
            ws1.cell(row, 2, name)
            ws1.cell(row, 3, spec or "")
            ws1.cell(row, 4, unit or "")
            ws1.cell(row, 5, hansun or 0)
            ws1.cell(row, 6, shiji)
            ws1.cell(row, 7, peijian)
            if peijian < -0.5:
                ws1.cell(row, 8, "超领(超用)")
            elif peijian > 0.5:
                ws1.cell(row, 8, "有余量(可移交)")
            else:
                ws1.cell(row, 8, "持平(对得上)")
            style_cell(ws1, row, 8)
            row += 1

# 开关插座面板结算
ws_switch = wb_src["开关插座面板结算"]
for r in range(6, 39):
    name = ws_switch.cell(row=r, column=2).value
    shiji = ws_switch.cell(row=r, column=5).value
    hansun = ws_switch.cell(row=r, column=30).value
    spec = ws_switch.cell(row=r, column=3).value
    if name and isinstance(shiji, (int,float)):
        peijian = (hansun or 0) - shiji if isinstance(hansun, (int,float)) else None
        if peijian is not None:
            ws1.cell(row, 1, "开关插座面板")
            ws1.cell(row, 2, name)
            ws1.cell(row, 3, spec or "")
            ws1.cell(row, 4, "个")
            ws1.cell(row, 5, hansun or 0)
            ws1.cell(row, 6, shiji)
            ws1.cell(row, 7, peijian)
            if peijian < -0.5:
                ws1.cell(row, 8, "超领(超用)")
            elif peijian > 0.5:
                ws1.cell(row, 8, "有余量(可移交)")
            else:
                ws1.cell(row, 8, "持平(对得上)")
            style_cell(ws1, row, 8)
            row += 1

# 浴霸及凉霸
ws_h = wb_src["浴霸及凉霸结算"]
for r in range(5, 8):
    name = ws_h.cell(row=r, column=2).value
    shiji = ws_h.cell(row=r, column=5).value
    heji = ws_h.cell(row=r, column=13).value
    if name and isinstance(shiji, (int,float)) and isinstance(heji, (int,float)):
        peijian = heji - shiji
        ws1.cell(row, 1, "浴霸及凉霸")
        ws1.cell(row, 2, name)
        ws1.cell(row, 3, "")
        ws1.cell(row, 4, "只")
        ws1.cell(row, 5, heji)
        ws1.cell(row, 6, shiji)
        ws1.cell(row, 7, peijian)
        ws1.cell(row, 8, "持平(对得上)")
        style_cell(ws1, row, 8)
        row += 1

# 洁具(科勒)
ws_k = wb_src["洁具结算（科勒）"]
for r in range(5, 10):
    name = ws_k.cell(row=r, column=2).value
    shiji = ws_k.cell(row=r, column=5).value
    heji = ws_k.cell(row=r, column=13).value
    if name and isinstance(shiji, (int,float)) and isinstance(heji, (int,float)):
        peijian = heji - shiji
        ws1.cell(row, 1, "洁具(科勒)")
        ws1.cell(row, 2, name)
        ws1.cell(row, 3, "")
        ws1.cell(row, 4, "套")
        ws1.cell(row, 5, heji)
        ws1.cell(row, 6, shiji)
        ws1.cell(row, 7, peijian)
        ws1.cell(row, 8, "持平" if peijian == 0 else "有余量")
        style_cell(ws1, row, 8)
        row += 1

ws1.column_dimensions["A"].width = 16
ws1.column_dimensions["B"].width = 35
ws1.column_dimensions["C"].width = 35
ws1.column_dimensions["D"].width = 8
ws1.column_dimensions["E"].width = 18
ws1.column_dimensions["F"].width = 18
ws1.column_dimensions["G"].width = 18
ws1.column_dimensions["H"].width = 20

# Sheet 2: 对不上数量的明细
ws2 = wb.create_sheet("对不上数量的明细")
ws2.cell(1,1,"物业配品配件 - 数量对不上/异常的明细").font = title_font
ws2.merge_cells("A1:H1")
ws2.cell(2,1,"（超领 = 施工领用超过合同含损耗量需扣款；有余量 = 有剩余可移交物业的配件）").font = Font(size=10, italic=True, color="666666")
ws2.merge_cells("A2:H2")

headers2 = ["材料类别", "材料名称", "单位", "含损耗数量", "施工领用量", "物业配品配件", "问题类型", "说明"]
for c, h in enumerate(headers2, 1):
    ws2.cell(3, c, h)
style_header(ws2, 3, 8)

row2 = 4
for r in range(4, ws1.max_row + 1):
    issue = ws1.cell(r, 8).value
    if issue and ("超领" in str(issue) or "有余量" in str(issue)):
        for c in range(1, 8):
            ws2.cell(row2, c).value = ws1.cell(r, c).value
            ws2.cell(row2, c).border = thin_border
            ws2.cell(row2, c).alignment = Alignment(horizontal="center", vertical="center")
        peijian = ws1.cell(r, 7).value or 0
        if peijian < -0.5:
            ws2.cell(row2, 7, "超领(需扣款)")
            ws2.cell(row2, 8, "施工领用超过含损耗量，多用的需从结算扣回")
            for c in range(1,9): ws2.cell(row2,c).fill = red_fill
        else:
            ws2.cell(row2, 7, "有余量(可移交物业)")
            ws2.cell(row2, 8, "施工领用少于含损耗量，剩余配件应交物业")
            for c in range(1,9): ws2.cell(row2,c).fill = green_fill
        ws2.cell(row2, 7).border = thin_border
        ws2.cell(row2, 8).border = thin_border
        ws2.cell(row2, 7).alignment = Alignment(horizontal="center", vertical="center")
        ws2.cell(row2, 8).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        row2 += 1

ws2.column_dimensions["A"].width = 16
ws2.column_dimensions["B"].width = 35
ws2.column_dimensions["C"].width = 8
ws2.column_dimensions["D"].width = 14
ws2.column_dimensions["E"].width = 14
ws2.column_dimensions["F"].width = 14
ws2.column_dimensions["G"].width = 18
ws2.column_dimensions["H"].width = 45

# Sheet 3: 可移交物业配件清单
ws3 = wb.create_sheet("可移交物业配件清单")
ws3.cell(1,1,"可移交物业的配件清单（有余量项）").font = title_font
ws3.merge_cells("A1:G1")

headers3 = ["材料类别", "材料名称", "规格型号", "单位", "含损耗数量", "施工领用量", "可移交物业数量"]
for c, h in enumerate(headers3, 1):
    ws3.cell(3, c, h)
style_header(ws3, 3, 7)

row3 = 4
for r in range(4, ws1.max_row + 1):
    peijian = ws1.cell(r, 7).value
    if isinstance(peijian, (int,float)) and peijian > 0.5:
        for c in range(1,8):
            ws3.cell(row3, c).value = ws1.cell(r, c).value
            ws3.cell(row3, c).border = thin_border
            ws3.cell(row3, c).alignment = Alignment(horizontal="center", vertical="center")
        row3 += 1

ws3.column_dimensions["A"].width = 16
ws3.column_dimensions["B"].width = 35
ws3.column_dimensions["C"].width = 35
ws3.column_dimensions["D"].width = 8
ws3.column_dimensions["E"].width = 16
ws3.column_dimensions["F"].width = 16
ws3.column_dimensions["G"].width = 18

out = r"D:\27968\我的文档\Users\27968\Documents\青羊58施工计划项目\上海远香湖_物业配件移交对比表.xlsx"
wb.save(out)
print("OK: " + out)
