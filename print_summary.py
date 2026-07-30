# -*- coding: utf-8 -*-
import io, sys
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
import openpyxl

wb = openpyxl.load_workbook("D:/27968/我的文档/Users/27968/Documents/青羊58施工计划项目/上海远香湖_物业备品配件数据说明.xlsx")
ws = wb["配品数据总说明"]

less_items = []
over_items = []
for r in range(6, ws.max_row + 1):
    ch = ws.cell(r, 9).value
    name = ws.cell(r, 2).value
    ht = ws.cell(r, 5).value
    pj = ws.cell(r, 6).value
    yl = ws.cell(r, 7).value
    sj = ws.cell(r, 8).value
    amt = ws.cell(r, 10).value
    dj = ws.cell(r, 12).value
    if isinstance(ch, (int,float)) and isinstance(ht, (int,float)):
        if ch > 0:
            less_items.append((name, ht, pj, yl, sj, ch, dj, amt))
        elif ch < 0:
            over_items.append((name, ht, pj, yl, sj, ch, dj, amt))

less_qty = sum(i[5] for i in less_items)
less_amt = sum(i[7] for i in less_items)
over_qty = sum(abs(i[5]) for i in over_items)
over_amt = sum(i[7] for i in over_items)

print("=" * 70)
print("少领项汇总(实际领用 < 应领总量)")
print("=" * 70)
less_printed = set()
for i in less_items:
    name, ht, pj, yl, sj, ch, dj, amt = i
    pj_s = f"+{pj:.1f}" if pj and pj > 0 else ""
    print(f"  {name:.<35}")
    print(f"    合同={ht:.0f}  {pj_s} 应领={yl:.0f} 实领={sj:.0f}  少领={ch:.0f}  @{dj:.2f}  ={amt:.2f}")

print()
print(f"少领合计: {len(less_items)}项, {less_qty:.0f}个/件, 金额={less_amt:.2f}")
print()

print("=" * 70)
print("超领项汇总(实际领用 > 应领总量)")
print("=" * 70)
for i in over_items[:15]:
    name, ht, pj, yl, sj, ch, dj, amt = i
    pj_s = f"+{pj:.1f}" if pj and pj > 0 else ""
    print(f"  {name:.<35}")
    print(f"    合同={ht:.0f}  {pj_s} 应领={yl:.0f} 实领={sj:.0f}  超领={abs(ch):.0f}  @{dj:.2f}  ={amt:.2f}")

if len(over_items) > 15:
    print(f"  ...还有{len(over_items)-15}项")
print()
print(f"超领合计: {len(over_items)}项, {over_qty:.0f}个/件, 金额={over_amt:.2f}")
