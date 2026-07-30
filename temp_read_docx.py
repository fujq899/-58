import pythoncom
pythoncom.CoInitialize()
from docx import Document
doc = Document("D:/形象进度(1).docx")
print("=== PARAGRAPHS ===")
for i, p in enumerate(doc.paragraphs):
    t = p.text.strip()
    if t:
        print(f"[P{i}] {t}")
print("=== TABLES ===")
for ti, table in enumerate(doc.tables):
    print(f"=== Table {ti} ===")
    for ri, row in enumerate(table.rows):
        cells = [c.text.strip() for c in row.cells]
        print(f"  Row {ri}: {' | '.join(cells)}")
print("=== DONE ===")
