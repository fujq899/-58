# -*- coding: utf-8 -*-
import openpyxl

out_path = "D:/27968/我的文档/Users/27968/Documents/青羊58施工计划项目/上海远香湖_物业备品配件数据说明.xlsx"
src_path = "D:/项目综合资料/天府办公文件/上海远香湖结算单/最总结算表格/02 结算前扣款事项汇总.xlsx"

wb = openpyxl.load_workbook(out_path)
wb_src = openpyxl.load_workbook(src_path, data_only=True)

ws = wb["配品数据总说明"]

# 构建源数据查询表
# 灯具: name→(hetong=col30, hansun=col31)
light_data = {}
for r in range(7, 36):
    n = wb_src["灯具结算"].cell(r, 2).value
    c30 = wb_src["灯具结算"].cell(r, 30).value
    c31 = wb_src["灯具结算"].cell(r, 31).value
    price = wb_src["灯具结算"].cell(r, 6).value
    unit = wb_src["灯具结算"].cell(r, 4).value
    if n and isinstance(c30, (int,float)):
        light_data[n] = {"hetong": c30, "hansun": c31 or c30, "price": price, "unit": unit}

# 开关面板: name→(hetong=col29, hansun=col30)
sw_data = {}
for r in range(6, 39):
    n = wb_src["开关插座面板结算"].cell(r, 2).value
    c29 = wb_src["开关插座面板结算"].cell(r, 29).value
    c30 = wb_src["开关插座面板结算"].cell(r, 30).value
    price = wb_src["开关插座面板结算"].cell(r, 6).value
    if n and isinstance(c29, (int,float)):
        sw_data[n] = {"hetong": c29, "hansun": c30 or c29, "price": price, "unit": "个"}

# HDPE管材: name→(hetong=col15, hansun=col20)
hdpe_data = {}
for r in range(5, 51):
    n = wb_src["HDPE管材结算"].cell(r, 3).value
    c15 = wb_src["HDPE管材结算"].cell(r, 15).value
    c20 = wb_src["HDPE管材结算"].cell(r, 20).value
    price = wb_src["HDPE管材结算"].cell(r, 9).value
    unit = wb_src["HDPE管材结算"].cell(r, 6).value
    if n and isinstance(c20, (int,float)):
        hetong = c15 if isinstance(c15, (int,float)) else 0
        hdpe_data[n] = {"hetong": hetong, "hansun": c20, "price": price, "unit": unit}

# 洁具科勒: name→(hetong=col13, hansun=col13)
kohler_data = {}
for r in range(5, 10):
    n = wb_src["洁具结算（科勒）"].cell(r, 2).value
    c13 = wb_src["洁具结算（科勒）"].cell(r, 13).value
    price = wb_src["洁具结算（科勒）"].cell(r, 6).value
    if n and isinstance(c13, (int,float)):
        kohler_data[n] = {"hetong": c13, "hansun": c13, "price": price, "unit": "套"}

def find_data(cat, name):
    if cat == "灯具":
        return light_data.get(name)
    elif cat == "开关插座面板":
        return sw_data.get(name)
    elif "HDPE" in str(cat):
        return hdpe_data.get(name)
    elif "洁具" in str(cat):
        return kohler_data.get(name)
    return None

# 填充合同量和配品列
filled = 0
for r in range(6, ws.max_row + 1):
    cat = str(ws.cell(r, 1).value or "")
    name = str(ws.cell(r, 2).value or "")
    d = find_data(cat, name)
    
    # 也尝试不带编号的匹配（因为xlsx里有编号）
    if not d:
        for k, v in (light_data if cat == "灯具" else sw_data if cat == "开关插座面板" else {}).items():
            if k and name and (k in name or name in k):
                d = v
                break
    
    if d:
        hetong = d["hetong"]
        hansun = d["hansun"]
        peijian = round(hansun - hetong, 2) if isinstance(hansun, (int,float)) and isinstance(hetong, (int,float)) else 0
        
        ws.cell(r, 5, hetong)  # E列: 合同清单量
        ws.cell(r, 6, peijian)  # F列: 物业配品配件
        filled += 1

print(f"已填充{filled}行的合同量和配品数据")

# 更新HDPE行（匹配规格型号）
for r in range(6, ws.max_row + 1):
    cat = str(ws.cell(r, 1).value or "")
    if "HDPE" in cat:
        spec = str(ws.cell(r, 3).value or "")
        for n, d in hdpe_data.items():
            spec2 = wb_src["HDPE管材结算"].cell(list(hdpe_data.keys()).index(n) + 5, 4).value or ""
            if spec and spec2 and (spec in spec2 or spec2 in spec):
                ws.cell(r, 5, d["hetong"])
                ws.cell(r, 6, round(d["hansun"] - d["hetong"], 2))
                filled += 1
                break

wb.save(out_path)
print("文件已更新保存")
