# -*- coding: utf-8 -*-
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

out_path = "D:/27968/我的文档/Users/27968/Documents/青羊58施工计划项目/上海远香湖_物业备品配件数据说明.xlsx"
src_path = "D:/项目综合资料/天府办公文件/上海远香湖结算单/最总结算表格/02 结算前扣款事项汇总.xlsx"

wb = openpyxl.load_workbook(out_path)
wb_src = openpyxl.load_workbook(src_path, data_only=True)

hdr_fill = PatternFill("solid", fgColor="4472C4")
hdr_font = Font(bold=True, size=10, color="FFFFFF")
thin = Border(left=Side("thin"), right=Side("thin"), top=Side("thin"), bottom=Side("thin"))
green_fill = PatternFill("solid", fgColor="C6EFCE")
red_fill = PatternFill("solid", fgColor="FFC7CE")

# 灯具单价
light_prices = {}
for r in range(7, 36):
    n = wb_src["灯具结算"].cell(r, 2).value
    p = wb_src["灯具结算"].cell(r, 6).value
    if n and isinstance(p, (int,float)): light_prices[n] = p

# 开关面板单价
sw_prices = {}
for r in range(6, 39):
    n = wb_src["开关插座面板结算"].cell(r, 2).value
    p = wb_src["开关插座面板结算"].cell(r, 6).value
    if n and isinstance(p, (int,float)): sw_prices[n] = p

# HDPE单价
hdpe_prices = {}
for r in range(5, 51):
    n = wb_src["HDPE管材结算"].cell(r, 3).value
    p = wb_src["HDPE管材结算"].cell(r, 9).value
    if n and isinstance(p, (int,float)): hdpe_prices[n] = p

def get_price(cat, name):
    if cat == "灯具":
        p = light_prices.get(name)
        if p: return p
        if "LED线条灯" in str(name): return 121.26
        if "灯带驱动" in str(name): return 74.70
        return None
    elif cat == "开关插座面板":
        return sw_prices.get(name)
    elif "HDPE" in str(cat):
        return hdpe_prices.get(name)
    return None

ws1 = wb["配品数据总说明"]

# 在最后添加两列：单价 + 数量说明
ws1.cell(5, 12, "单价(含税)").font = hdr_font
ws1.cell(5, 12).fill = hdr_fill
ws1.cell(5, 12).border = thin
ws1.cell(5, 12).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

ws1.cell(5, 13, "数量明细说明").font = hdr_font
ws1.cell(5, 13).fill = hdr_fill
ws1.cell(5, 13).border = thin
ws1.cell(5, 13).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

for r in range(6, ws1.max_row + 1):
    cat = str(ws1.cell(r, 1).value or "")
    name = str(ws1.cell(r, 2).value or "")
    unit = str(ws1.cell(r, 4).value or "")
    hetong = ws1.cell(r, 5).value
    peijian = ws1.cell(r, 6).value
    yingling = ws1.cell(r, 7).value
    shiji = ws1.cell(r, 8).value
    chae = ws1.cell(r, 9).value
    
    price = get_price(cat, name)
    
    # 单价
    if price:
        ws1.cell(r, 12, price)
    else:
        ws1.cell(r, 12, "")
    
    # 数量明细说明
    if isinstance(chae, (int,float)) and isinstance(yingling, (int,float)) and isinstance(shiji, (int,float)):
        if chae > 0:
            note = "少领%(gap).0f%(u)s：应领%(yl).0f%(u)s - 实际领用%(sj).0f%(u)s" % {"gap": chae, "u": unit, "yl": yingling, "sj": shiji}
        elif chae < 0:
            note = "超领%(gap).0f%(u)s：实际领用%(sj).0f%(u)s - 应领%(yl).0f%(u)s" % {"gap": abs(chae), "u": unit, "yl": yingling, "sj": shiji}
        else:
            note = "持平：应领%(yl).0f%(u)s = 实际领用%(sj).0f%(u)s" % {"yl": yingling, "sj": shiji, "u": unit}
        ws1.cell(r, 13, note)
    
    for c in [12, 13]:
        ws1.cell(r, c).border = thin
        ws1.cell(r, c).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

ws1.column_dimensions["L"].width = 14
ws1.column_dimensions["M"].width = 45

# 同样更新少领明细sheet
ws3 = wb["少领明细(可移交物业)"]
ws3.cell(3, 9, "单价(含税)").font = hdr_font
ws3.cell(3, 9).fill = hdr_fill
ws3.cell(3, 9).border = thin
ws3.cell(3, 9).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
ws3.cell(3, 10, "数量明细说明").font = hdr_font
ws3.cell(3, 10).fill = hdr_fill
ws3.cell(3, 10).border = thin
ws3.cell(3, 10).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

for r in range(4, ws3.max_row + 1):
    cat = str(ws3.cell(r, 1).value or "")
    name = str(ws3.cell(r, 2).value or "")
    unit = str(ws3.cell(r, 3).value or "")
    yingling = ws3.cell(r, 4).value
    shiji = ws3.cell(r, 5).value
    chae_val = ws3.cell(r, 6).value
    
    price = get_price(cat, name)
    if price:
        ws3.cell(r, 9, price)
    else:
        ws3.cell(r, 9, "")
    
    if isinstance(chae_val, (int,float)) and isinstance(yingling, (int,float)) and isinstance(shiji, (int,float)):
        note = "应领%(yl).0f%(u)s - 实领%(sj).0f%(u)s = 少%(gap).0f%(u)s" % {"yl": yingling, "sj": shiji, "gap": chae_val, "u": unit}
        ws3.cell(r, 10, note)
    
    for c in [9, 10]:
        ws3.cell(r, c).border = thin
        ws3.cell(r, c).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

ws3.column_dimensions["I"].width = 14
ws3.column_dimensions["J"].width = 45

# 同样更新超领明细sheet
ws4 = wb["超领明细(需扣款)"]
ws4.cell(3, 9, "单价(含税)").font = hdr_font
ws4.cell(3, 9).fill = hdr_fill
ws4.cell(3, 9).border = thin
ws4.cell(3, 9).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
ws4.cell(3, 10, "数量明细说明").font = hdr_font
ws4.cell(3, 10).fill = hdr_fill
ws4.cell(3, 10).border = thin
ws4.cell(3, 10).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

for r in range(4, ws4.max_row + 1):
    cat = str(ws4.cell(r, 1).value or "")
    name = str(ws4.cell(r, 2).value or "")
    unit = str(ws4.cell(r, 3).value or "")
    yingling = ws4.cell(r, 4).value
    shiji = ws4.cell(r, 5).value
    chae_val = ws4.cell(r, 6).value
    
    price = get_price(cat, name)
    if price:
        ws4.cell(r, 9, price)
    else:
        ws4.cell(r, 9, "")
    
    if isinstance(chae_val, (int,float)) and isinstance(yingling, (int,float)) and isinstance(shiji, (int,float)):
        note = "实领%(sj).0f%(u)s - 应领%(yl).0f%(u)s = 超%(gap).0f%(u)s" % {"yl": yingling, "sj": shiji, "gap": abs(chae_val), "u": unit}
        ws4.cell(r, 10, note)
    
    for c in [9, 10]:
        ws4.cell(r, c).border = thin
        ws4.cell(r, c).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

ws4.column_dimensions["I"].width = 14
ws4.column_dimensions["J"].width = 45

wb.save(out_path)
print("OK")
