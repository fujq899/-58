# -*- coding: utf-8 -*-
import openpyxl

out_path = r"D:\27968\我的文档\Users\27968\Documents\青羊58施工计划项目\上海远香湖_物业配件移交对比表.xlsx"
src_path = r"D:\项目综合资料\天府办公文件\上海远香湖结算单\最总结算表格\02 结算前扣款事项汇总.xlsx"

wb_out = openpyxl.load_workbook(out_path)
ws1 = wb_out["物业配品配件对比总表"]
wb_src = openpyxl.load_workbook(src_path, data_only=True)

light_prices = {}
ws_light = wb_src["灯具结算"]
for r in range(7, 36):
    name = ws_light.cell(row=r, column=2).value
    price = ws_light.cell(row=r, column=6).value
    if name and isinstance(price, (int,float)):
        light_prices[name] = price

switch_prices = {}
ws_switch = wb_src["开关插座面板结算"]
for r in range(6, 39):
    name = ws_switch.cell(row=r, column=2).value
    price = ws_switch.cell(row=r, column=6).value
    if name and isinstance(price, (int,float)):
        switch_prices[name] = price

print("=" * 70)
print("物业移交对比表 - 超领项金额汇总")
print("（施工领用超过合同含损耗量的项）")
print("=" * 70)

total_chao = 0
total_amount = 0
results = []

for r in range(4, ws1.max_row + 1):
    issue = ws1.cell(r, 8).value
    if issue and "超领" in str(issue):
        cat = str(ws1.cell(r, 1).value or "")
        name = str(ws1.cell(r, 2).value or "")
        unit = str(ws1.cell(r, 4).value or "")
        hansun = ws1.cell(r, 5).value or 0
        shiji = ws1.cell(r, 6).value or 0
        peijian = ws1.cell(r, 7).value or 0
        chao = abs(peijian)

        price = None
        if cat == "灯具":
            price = light_prices.get(name)
        elif cat == "开关插座面板":
            price = switch_prices.get(name)

        if price:
            amount = chao * price
            total_chao += chao
            total_amount += amount
            results.append((cat, name, unit, hansun, shiji, chao, price, amount))
            print()
            print("[%s] %s" % (cat, name))
            print("  合同含损耗: %.1f  施工领用: %.0f  超领: %.0f%s" % (hansun, shiji, chao, unit))
            print("  单价: %.4f  超领金额: %.2f" % (price, amount))

print()
print("=" * 70)
print("超领项合计: 共%d项, 超领总数量: %.0f" % (len(results), total_chao))
print("超领总金额（含税）: %.2f" % total_amount)

by_cat = {}
for r in results:
    cat = r[0]
    if cat not in by_cat:
        by_cat[cat] = {"qty": 0, "amt": 0}
    by_cat[cat]["qty"] += r[5]
    by_cat[cat]["amt"] += r[7]

print()
print("按材料类别汇总：")
for cat, data in sorted(by_cat.items(), key=lambda x: -x[1]["amt"]):
    print("  %s: 超领数量 %.0f个, 超领金额 %.2f" % (cat, data["qty"], data["amt"]))
print("  总计: %.2f" % total_amount)

print()
print("=" * 70)
print("有余量项（可移交物业）汇总")
print("=" * 70)

total_yuliang = 0
total_spare_value = 0
for r in range(4, ws1.max_row + 1):
    issue = ws1.cell(r, 8).value
    if issue and "有余量" in str(issue):
        cat = str(ws1.cell(r, 1).value or "")
        name = str(ws1.cell(r, 2).value or "")
        unit = str(ws1.cell(r, 4).value or "")
        peijian = ws1.cell(r, 7).value or 0
        total_yuliang += peijian

        price = None
        if cat == "灯具":
            price = light_prices.get(name)
        elif cat == "开关插座面板":
            price = switch_prices.get(name)
        if price:
            val = peijian * price
            total_spare_value += val
            print("  %s: 余%.0f%s x %.2f = %.2f" % (name, peijian, unit, price, val))
        else:
            print("  %s: 余%.0f%s" % (name, peijian, unit))

print()
print("有余量合计: %.0f个/件可移交物业" % total_yuliang)
print("余量材料价值约: %.2f" % total_spare_value)

print()
print("=" * 70)
print("全部甲供材超领扣款汇总")
print("=" * 70)

ws_sum = wb_src["甲供材超供汇总表"]
for r in range(4, 21):
    name = ws_sum.cell(row=r, column=2).value
    hanhan = ws_sum.cell(row=r, column=3).value
    hanhan_han = ws_sum.cell(row=r, column=5).value
    if name and isinstance(hanhan, (int,float)) and hanhan != 0:
        print("  %s: 除税 %.2f  含税 %.2f" % (name, hanhan, hanhan_han))
